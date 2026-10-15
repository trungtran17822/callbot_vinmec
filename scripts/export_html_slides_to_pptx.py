from __future__ import annotations

import base64
import re
import shutil
import struct
import subprocess
import tempfile
import time
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
HTML = ROOT / "html-slides" / "index.html"
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
OUTPUT = ROOT / "Callbot_Vinmec_17_slides.pptx"
WIDTH, HEIGHT = 1600, 900
EMU_W, EMU_H = 12192000, 6858000


def png_size(path: Path) -> tuple[int, int]:
    with path.open("rb") as stream:
        header = stream.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"Edge did not produce a PNG: {path}")
    return struct.unpack(">II", header[16:24])


def capture_slides() -> list[Path]:
    source = HTML.read_text(encoding="utf-8")
    count = len(re.findall(r"<section\b[^>]*\bclass=[\"'][^\"']*\bslide\b", source))
    if count != 17:
        raise ValueError(f"Expected 17 HTML slides, found {count}")

    work = ROOT / ".tmp-pptx-export"
    if work.exists():
        shutil.rmtree(work, ignore_errors=True)
    work.mkdir()
    profile = work / "edge-profile"
    captures: list[Path] = []
    try:
        for index in range(1, count + 1):
            target = work / f"slide-{index:02d}.png"
            url = HTML.as_uri() + f"#{index}"
            command = [
                str(EDGE),
                "--headless",
                "--disable-gpu",
                "--disable-crash-reporter",
                "--hide-scrollbars",
                "--allow-file-access-from-files",
                "--no-first-run",
                f"--user-data-dir={profile}",
                f"--window-size={WIDTH},{HEIGHT}",
                f"--screenshot={target}",
                url,
            ]
            result = subprocess.run(command, check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            for _ in range(30):
                if target.exists() and target.stat().st_size > 10_000:
                    break
                time.sleep(0.2)
            if result.returncode or not target.exists() or png_size(target) != (WIDTH, HEIGHT):
                raise RuntimeError(f"Slide {index} capture failed (Edge exit {result.returncode})")
            captures.append(target)
        write_presentation(captures, OUTPUT)
    finally:
        subprocess.run(
            ["taskkill", "/F", "/IM", "msedge.exe", "/T"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            check=False,
        )
        shutil.rmtree(work, ignore_errors=True)
    return captures


def write_presentation(images: list[Path], output: Path) -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]
    for image in images:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(output)


if __name__ == "__main__":
    capture_slides()
    deck = Presentation(OUTPUT)
    if len(deck.slides) != 17:
        raise RuntimeError(f"Expected 17 slides in PPTX, got {len(deck.slides)}")
    with zipfile.ZipFile(OUTPUT) as archive:
        for member in archive.namelist():
            if member.endswith(".xml") or member.endswith(".rels"):
                ET.fromstring(archive.read(member))
        if archive.testzip() is not None:
            raise RuntimeError("PPTX zip integrity check failed")
    print(f"Created {OUTPUT} ({OUTPUT.stat().st_size:,} bytes, {len(deck.slides)} slides)")
